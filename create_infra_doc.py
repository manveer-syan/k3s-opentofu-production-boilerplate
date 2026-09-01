import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_document():
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("ATE OPERATIONS PLATFORM")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Infrastructure & Deployment Plan (AWS, Terraform, Docker, K8s, GitLab CI)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Metadata Box Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    metadata_pairs = [
        ("Project Name:", "ATE Operations Management Suite (Golang Backend)"),
        ("Infrastructure Provider:", "Amazon Web Services (AWS) - $100 Credit / Free Tier"),
        ("DevOps Stack:", "Terraform IaC, Docker, Nginx, Ansible, Kubernetes, GitLab CI"),
        ("Document Version:", "1.0 (Production Release)")
    ]

    for i, (k, v) in enumerate(metadata_pairs):
        row = meta_table.rows[i]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.3)
        
        set_cell_background(cell_k, "F8FAFC")
        set_cell_background(cell_v, "F1F5F9")
        
        r_k = cell_k.paragraphs[0].add_run(k)
        r_k.font.bold = True
        r_k.font.size = Pt(10)
        r_k.font.name = "Arial"
        
        r_v = cell_v.paragraphs[0].add_run(v)
        r_v.font.size = Pt(10)
        r_v.font.name = "Arial"

    doc.add_paragraph()

    # Section 1: Executive Summary & AWS Cost Strategy
    h1 = doc.add_heading("1. Executive Summary & Cost Strategy", level=1)
    h1.style.font.name = "Arial"
    h1.style.font.color.rgb = RGBColor(15, 23, 42)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(
        "This document defines the end-to-end Infrastructure as Code (IaC) deployment plan for the ATE Operations Platform. "
        "The application is engineered as a lightweight, high-performance Golang microservice with an embedded monochrome Web UI. "
        "The infrastructure is designed to run seamlessly within the AWS Free Tier ($100 promotional credits), utilizing automated "
        "provisioning via Terraform, containerization via Docker, configuration management via Ansible, and CI/CD automation via GitLab CI."
    )
    r.font.name = "Arial"
    r.font.size = Pt(10.5)

    doc.add_paragraph()

    # Section 2: Architecture Overview
    h2 = doc.add_heading("2. System Architecture Diagram", level=1)
    h2.style.font.name = "Arial"
    h2.style.font.color.rgb = RGBColor(15, 23, 42)

    arch_text = """+-----------------------------------------------------------------------+
|                             CLIENT / BROWSER                          |
|                    (Minimal Monochrome Technical UI)                  |
+-----------------------------------------------------------------------+
                                   |
                         HTTP Port 8080 (REST API / UI)
                                   v
+-----------------------------------------------------------------------+
|                    AWS EC2 INSTANCE (t3.micro / Ubuntu)               |
|                                                                       |
|   +---------------------------------------------------------------+   |
|   |                  DOCKER ENGINE & CONTAINER                     |   |
|   |                  (Go Microservice - ~15MB)                     |   |
|   |                                                               |   |
|   |   +-------------------------+    +------------------------+   |   |
|   |   | REST API & Health Check |    | Thread-Safe Store      |   |   |
|   |   | (/api/v1/operations)    |    | (sync.RWMutex Engine)  |   |   |
|   |   +-------------------------+    +------------------------+   |   |
|   +---------------------------------------------------------------+   |
+-----------------------------------------------------------------------+
                                   |
       +---------------------------+---------------------------+
       |                           |                           |
       v                           v                           v
+--------------+           +--------------+           +------------------+
|  TERRAFORM   |           |   ANSIBLE    |           |    KUBERNETES    |
| (AWS EC2/EIP)|           | (Playbook)   |           | (Deployment/HPA) |
+--------------+           +--------------+           +------------------+"""

    arch_p = doc.add_paragraph()
    r_arch = arch_p.add_run(arch_text)
    r_arch.font.name = "Courier New"
    r_arch.font.size = Pt(8.5)

    doc.add_paragraph()

    # Section 3: Infrastructure Breakdown
    h3 = doc.add_heading("3. Infrastructure Components Breakdown", level=1)
    h3.style.font.name = "Arial"
    h3.style.font.color.rgb = RGBColor(15, 23, 42)

    components = [
        ("AWS EC2 Instance (Compute)", "Provisions an Ubuntu 22.04 LTS t3.micro server instance. Eligible for 750 free hours/month under AWS Free Tier."),
        ("AWS Security Group (Firewall)", "Configures stateful firewall rules: Port 80 (HTTP), Port 443 (HTTPS), Port 8080 (Go Application), and Port 22 (SSH Admin)."),
        ("AWS Elastic IP (Static Networking)", "Assigns a permanent, public-facing IPv4 address to ensure zero DNS breakage across instance restarts."),
        ("Terraform IaC (terraform/aws_main.tf)", "Declarative automation script that builds AWS VPC, Security Group, EC2 server, and bootstraps Docker automatically."),
        ("Docker Container Engine", "Runs the compiled multi-stage Go binary inside an Alpine Linux container (~15MB size) with auto-restart policies."),
        ("Ansible Playbook (ansible/playbook.yml)", "Configuration management playbook for installing packages, pulling code, and verifying /health API response."),
        ("Kubernetes Manifests (k8s/)", "Enterprise manifests including Deployment (3 replicas), NodePort Service, Ingress routing, and Horizontal Pod Autoscaler (HPA)."),
        ("GitLab CI/CD (.gitlab-ci.yml)", "Automated pipeline with 4 stages: unit_tests, compile_binary, docker_build, and deploy_staging.")
    ]

    for title, desc in components:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.2)
        r_t = p_c.add_run(f"• {title}: ")
        r_t.font.bold = True
        r_t.font.size = Pt(10)
        r_t.font.name = "Arial"
        
        r_d = p_c.add_run(desc)
        r_d.font.size = Pt(10)
        r_d.font.name = "Arial"

    doc.add_paragraph()

    # Section 4: Step-by-Step Deployment Guide
    h4 = doc.add_heading("4. Step-by-Step AWS Deployment Execution Guide", level=1)
    h4.style.font.name = "Arial"
    h4.style.font.color.rgb = RGBColor(15, 23, 42)

    steps = [
        ("Step 1: Configure AWS CLI Credentials", 
         "Obtain your AWS Access Key ID and Secret Access Key from AWS Console > Security Credentials.\nSet environment variables on your workstation:\nexport AWS_ACCESS_KEY_ID=\"your_access_key\"\nexport AWS_SECRET_ACCESS_KEY=\"your_secret_key\"\nexport AWS_DEFAULT_REGION=\"us-east-1\""),

        ("Step 2: Deploy Infrastructure via Terraform", 
         "Navigate to the terraform directory and execute the Terraform workflow:\ncd terraform\nterraform init\nterraform apply -auto-approve\n\nTerraform will output your public server IP and application URL upon completion."),

        ("Step 3: Verification & Health Check", 
         "Verify your live AWS deployment by testing the healthcheck endpoint:\ncurl http://<YOUR_AWS_PUBLIC_IP>:8080/health\n\nExpected JSON Response:\n{\"service\":\"ate-operations-go\",\"status\":\"healthy\",\"timestamp\":\"...\",\"uptime\":\"running\"}"),

        ("Step 4: Execute Ansible Configuration Management", 
         "To run automated configuration updates across your servers:\ncd ansible\nansible-playbook -i inventory.ini playbook.yml"),

        ("Step 5: Apply Kubernetes Manifests (Optional K8s Cluster)", 
         "If deploying to an AWS EKS or Kubernetes cluster:\nkubectl apply -f k8s/")
    ]

    for title, detail in steps:
        p_step = doc.add_paragraph()
        r_st = p_step.add_run(title)
        r_st.font.bold = True
        r_st.font.size = Pt(11)
        r_st.font.name = "Arial"
        r_st.font.color.rgb = RGBColor(30, 41, 59)

        p_det = doc.add_paragraph()
        p_det.paragraph_format.left_indent = Inches(0.3)
        r_dt = p_det.add_run(detail)
        r_dt.font.size = Pt(9.5)
        r_dt.font.name = "Courier New" if "export" in detail or "terraform" in detail or "curl" in detail else "Arial"

    doc.add_paragraph()

    # Section 5: Operations Reference
    h5 = doc.add_heading("5. DevOps Makefile Reference Matrix", level=1)
    h5.style.font.name = "Arial"
    h5.style.font.color.rgb = RGBColor(15, 23, 42)

    make_table = doc.add_table(rows=10, cols=2)
    make_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = [("Makefile Command", "Operational Description")]
    make_commands = [
        ("make build", "Compiles Go binary executable (ate-app)"),
        ("make run", "Compiles and executes local Go server"),
        ("make test", "Runs Go unit test suite (store/store_test.go)"),
        ("make docker-build", "Builds multi-stage Docker image (~15MB size)"),
        ("make tf-init", "Initializes Terraform working directory"),
        ("make tf-apply", "Provisions AWS EC2 & Elastic IP infrastructure"),
        ("make ansible-play", "Executes Ansible playbook configuration"),
        ("make k8s-apply", "Applies Kubernetes manifests to cluster"),
        ("make clean", "Removes compiled binary artifacts")
    ]

    for i, (cmd, desc) in enumerate(headers + make_commands):
        row = make_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.3)
        
        if i == 0:
            set_cell_background(c0, "0F172A")
            set_cell_background(c1, "0F172A")
            r0 = c0.paragraphs[0].add_run(cmd)
            r0.font.bold = True
            r0.font.color.rgb = RGBColor(255, 255, 255)
            r0.font.size = Pt(10)
            
            r1 = c1.paragraphs[0].add_run(desc)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(255, 255, 255)
            r1.font.size = Pt(10)
        else:
            bg_color = "F8FAFC" if i % 2 == 1 else "FFFFFF"
            set_cell_background(c0, bg_color)
            set_cell_background(c1, bg_color)
            
            r0 = c0.paragraphs[0].add_run(cmd)
            r0.font.name = "Courier New"
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            
            r1 = c1.paragraphs[0].add_run(desc)
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)

    doc.save("ATE_Infrastructure_Plan_and_Steps.docx")
    print("Successfully created ATE_Infrastructure_Plan_and_Steps.docx")

if __name__ == "__main__":
    create_document()
